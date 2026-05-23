import os
import json
import base64
import io
import zipfile
from contextlib import asynccontextmanager
from typing import List

import httpx
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from openai import AsyncOpenAI
from pydantic import BaseModel

load_dotenv()

DASHSCOPE_API_KEY = os.getenv("DASHSCOPE_API_KEY", "")
DASHSCOPE_BASE_URL = "https://dashscope-intl.aliyuncs.com/compatible-mode/v1"
DASHSCOPE_NATIVE_URL = "https://dashscope-intl.aliyuncs.com/api/v1"
FREEMODEL_API_KEY = os.getenv("FREEMODEL_API_KEY", "")
FREEMODEL_BASE_URL = "https://api.freemodel.dev/v1"
NVIDIA_API_KEY = os.getenv("NVIDIA_API_KEY", "")
NVIDIA_BASE_URL = "https://integrate.api.nvidia.com/v1"

TEXT_MODEL_PROVIDERS = {
    "qwen-plus": "dashscope",
    "qwen-max": "dashscope",
    "qwen-turbo": "dashscope",
    "qwen3-235b-a22b": "dashscope",
    "qwen3-32b": "dashscope",
    "qwen3-14b": "dashscope",
    "qwq-plus": "dashscope",
    "qwen3-coder-plus": "dashscope",
    "gpt-5.5": "freemodel",
    "gpt-5.4": "freemodel",
    "gpt-5.4-mini": "freemodel",
    "gpt-5.3-codex": "freemodel",
    "meta/llama-3.3-70b-instruct": "nvidia",
    "meta/llama-3.1-70b-instruct": "nvidia",
    "nvidia/llama-3.1-nemotron-nano-8b-v1": "nvidia",
    "nvidia/llama-3.1-nemotron-51b-instruct": "nvidia",
}

client: AsyncOpenAI | None = None
freemodel_client: AsyncOpenAI | None = None
nvidia_client: AsyncOpenAI | None = None
http_client: httpx.AsyncClient | None = None


@asynccontextmanager
async def lifespan(app: FastAPI):
    global client, freemodel_client, nvidia_client, http_client
    client = AsyncOpenAI(
        base_url=DASHSCOPE_BASE_URL,
        api_key=DASHSCOPE_API_KEY,
    )
    freemodel_client = AsyncOpenAI(
        base_url=FREEMODEL_BASE_URL,
        api_key=FREEMODEL_API_KEY,
    )
    nvidia_client = AsyncOpenAI(
        base_url=NVIDIA_BASE_URL,
        api_key=NVIDIA_API_KEY,
    )
    http_client = httpx.AsyncClient(timeout=120.0)
    yield
    await http_client.aclose()


app = FastAPI(title="AI Hub", version="2.0.0", lifespan=lifespan)


@app.middleware("http")
async def add_no_cache_headers(request, call_next):
    response = await call_next(request)
    if request.url.path == "/" or request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

AVAILABLE_MODELS = {
    "text": [
        {"id": "qwen-plus", "name": "Qwen Plus", "provider": "Alibaba Cloud"},
        {"id": "qwen-max", "name": "Qwen Max", "provider": "Alibaba Cloud"},
        {"id": "qwen-turbo", "name": "Qwen Turbo", "provider": "Alibaba Cloud"},
        {"id": "qwen3-235b-a22b", "name": "Qwen3 235B", "provider": "Alibaba Cloud"},
        {"id": "qwen3-32b", "name": "Qwen3 32B", "provider": "Alibaba Cloud"},
        {"id": "qwen3-14b", "name": "Qwen3 14B", "provider": "Alibaba Cloud"},
        {"id": "qwq-plus", "name": "QwQ Plus (Reasoning)", "provider": "Alibaba Cloud"},
        {"id": "qwen3-coder-plus", "name": "Qwen3 Coder Plus", "provider": "Alibaba Cloud"},
        {"id": "gpt-5.5", "name": "GPT-5.5", "provider": "FreeModel"},
        {"id": "gpt-5.4", "name": "GPT-5.4", "provider": "FreeModel"},
        {"id": "gpt-5.4-mini", "name": "GPT-5.4 Mini", "provider": "FreeModel"},
        {"id": "gpt-5.3-codex", "name": "GPT-5.3 Codex", "provider": "FreeModel"},
        {"id": "meta/llama-3.3-70b-instruct", "name": "Llama 3.3 70B Instruct", "provider": "NVIDIA NIM"},
        {"id": "meta/llama-3.1-70b-instruct", "name": "Llama 3.1 70B Instruct", "provider": "NVIDIA NIM"},
        {"id": "nvidia/llama-3.1-nemotron-nano-8b-v1", "name": "Nemotron Nano 8B", "provider": "NVIDIA NIM"},
        {"id": "nvidia/llama-3.1-nemotron-51b-instruct", "name": "Nemotron 51B Instruct", "provider": "NVIDIA NIM"},
    ],
    "image": [
        {"id": "wan2.6-t2i", "name": "Wan 2.6 Text-to-Image", "provider": "Alibaba Cloud"},
        {"id": "wan2.7-image-pro", "name": "Wan 2.7 Image Pro", "provider": "Alibaba Cloud"},
        {"id": "wan2.7-image", "name": "Wan 2.7 Image", "provider": "Alibaba Cloud"},
    ],
    "tts": [
        {"id": "cosyvoice-v3-flash", "name": "CosyVoice v3 Flash", "provider": "Alibaba Cloud"},
        {"id": "qwen3-tts-flash", "name": "Qwen3 TTS Flash", "provider": "Alibaba Cloud"},
    ],
}


class ChatRequest(BaseModel):
    message: str
    model: str = "qwen-plus"
    system_prompt: str = ""
    temperature: float = 0.7
    stream: bool = True
    api_key: str = ""


class ImageRequest(BaseModel):
    prompt: str
    model: str = "wan2.6-t2i"
    size: str = "1024x1024"
    quality: str = "standard"
    api_key: str = ""


class TTSRequest(BaseModel):
    text: str
    provider: str = "qwen3-tts-flash"
    voice: str = "Cherry"
    api_key: str = ""


class ContentRequest(BaseModel):
    topic: str
    content_type: str = "blog"
    model: str = "qwen-plus"
    tone: str = "professional"
    language: str = "id"
    api_key: str = ""


@app.get("/api/models")
async def get_models():
    return AVAILABLE_MODELS


def get_text_client(model: str, api_key: str = "") -> AsyncOpenAI:
    provider = TEXT_MODEL_PROVIDERS.get(model, "dashscope")
    if provider == "freemodel":
        key = api_key.strip() or FREEMODEL_API_KEY
        if not key:
            raise HTTPException(
                status_code=401,
                detail="FreeModel API key is missing. Paste it in API Management or add FREEMODEL_API_KEY to .env.",
            )
        if api_key.strip():
            return AsyncOpenAI(base_url=FREEMODEL_BASE_URL, api_key=key)
        if not freemodel_client:
            raise HTTPException(status_code=500, detail="FreeModel client not initialized")
        return freemodel_client

    if provider == "nvidia":
        key = api_key.strip() or NVIDIA_API_KEY
        if not key:
            raise HTTPException(
                status_code=401,
                detail="NVIDIA NIM API key is missing. Paste it in API Management or add NVIDIA_API_KEY to .env.",
            )
        if api_key.strip():
            return AsyncOpenAI(base_url=NVIDIA_BASE_URL, api_key=key)
        if not nvidia_client:
            raise HTTPException(status_code=500, detail="NVIDIA NIM client not initialized")
        return nvidia_client

    key = api_key.strip() or DASHSCOPE_API_KEY
    if not key:
        raise HTTPException(
            status_code=401,
            detail="DashScope API key is missing. Paste it in API Management or add DASHSCOPE_API_KEY to .env.",
        )
    if api_key.strip():
        return AsyncOpenAI(base_url=DASHSCOPE_BASE_URL, api_key=key)
    if not client:
        raise HTTPException(status_code=500, detail="AI client not initialized")
    return client


def get_dashscope_key(api_key: str = "") -> str:
    key = api_key.strip() or DASHSCOPE_API_KEY
    if not key:
        raise HTTPException(
            status_code=401,
            detail="DashScope API key is missing. Paste it in API Management or add DASHSCOPE_API_KEY to .env.",
        )
    return key


@app.post("/api/chat")
async def chat(req: ChatRequest):
    selected_client = get_text_client(req.model, req.api_key)

    messages = []
    if req.system_prompt:
        messages.append({"role": "system", "content": req.system_prompt})
    messages.append({"role": "user", "content": req.message})

    if req.stream:
        async def generate():
            stream = await selected_client.chat.completions.create(
                model=req.model,
                messages=messages,
                temperature=req.temperature,
                stream=True,
            )
            async for chunk in stream:
                if chunk.choices and chunk.choices[0].delta.content:
                    data = json.dumps({"content": chunk.choices[0].delta.content})
                    yield f"data: {data}\n\n"
            yield "data: [DONE]\n\n"

        return StreamingResponse(generate(), media_type="text/event-stream")

    response = await selected_client.chat.completions.create(
        model=req.model,
        messages=messages,
        temperature=req.temperature,
    )
    return {"content": response.choices[0].message.content}


@app.post("/api/image/generate")
async def generate_image(req: ImageRequest):
    if not http_client:
        raise HTTPException(status_code=500, detail="HTTP client not initialized")

    url = f"{DASHSCOPE_NATIVE_URL}/services/aigc/multimodal-generation/generation"
    dashscope_key = get_dashscope_key(req.api_key)
    headers = {
        "Authorization": f"Bearer {dashscope_key}",
        "Content-Type": "application/json",
    }

    payload = {
        "model": req.model,
        "input": {
            "messages": [
                {
                    "role": "user",
                    "content": [{"text": req.prompt}],
                }
            ]
        },
        "parameters": {
            "size": req.size,
            "n": 1,
            "watermark": False,
        },
    }

    try:
        resp = await http_client.post(url, json=payload, headers=headers)
        if resp.status_code != 200:
            raise HTTPException(status_code=resp.status_code, detail=resp.text)

        result = resp.json()
        output = result.get("output", {})
        choices = output.get("choices", [])

        if choices:
            content = choices[0].get("message", {}).get("content", [])
            for item in content:
                if "image" in item:
                    return {"url": item["image"]}

        raise HTTPException(status_code=500, detail="No image generated")
    except httpx.HTTPError as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/tts")
async def text_to_speech(req: TTSRequest):
    if not http_client:
        raise HTTPException(status_code=500, detail="HTTP client not initialized")

    url = f"{DASHSCOPE_NATIVE_URL}/services/aigc/multimodal-generation/generation"
    dashscope_key = get_dashscope_key(req.api_key)
    headers = {
        "Authorization": f"Bearer {dashscope_key}",
        "Content-Type": "application/json",
    }

    payload = {
        "model": req.provider,
        "input": {
            "text": req.text,
            "voice": req.voice,
        },
    }

    try:
        resp = await http_client.post(url, json=payload, headers=headers)
        if resp.status_code != 200:
            raise HTTPException(status_code=resp.status_code, detail=resp.text)

        result = resp.json()
        output = result.get("output", {})
        audio_data = output.get("audio", {})
        audio_url = audio_data.get("url", "") if isinstance(audio_data, dict) else audio_data

        if audio_url:
            audio_resp = await http_client.get(audio_url)
            if audio_resp.status_code != 200:
                raise HTTPException(status_code=502, detail="Failed to download generated audio")
            return StreamingResponse(
                iter([audio_resp.content]),
                media_type="audio/mpeg",
                headers={"Content-Disposition": "attachment; filename=speech.mp3"},
            )

        raise HTTPException(status_code=500, detail="No audio generated")
    except httpx.HTTPError as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/content/generate")
async def generate_content(req: ContentRequest):
    selected_client = get_text_client(req.model, req.api_key)

    prompts = {
        "blog": f"Write a comprehensive blog post about: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include a title, introduction, main sections with subheadings, and conclusion. Use markdown formatting.",
        "social": f"Create 5 engaging social media posts about: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include hashtags and emojis. Format each post clearly with a separator.",
        "script": f"Write a video script about: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include: Hook (first 3 seconds), Introduction, Main content with timestamps, Call to action, and Outro.",
        "email": f"Write a professional email about: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include subject line, greeting, body, and sign-off.",
        "ad": f"Create 3 advertising copy variations for: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include headline, body copy, and call-to-action for each.",
        "marketing": f"Create a comprehensive marketing plan outline for: {req.topic}. Tone: {req.tone}. Language: {req.language}. Include target audience, key messages, channels, and KPIs.",
    }

    system_prompt = "You are a professional content creator and copywriter. Create high-quality, engaging content."
    user_prompt = prompts.get(req.content_type, prompts["blog"])

    async def generate():
        stream = await selected_client.chat.completions.create(
            model=req.model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.8,
            stream=True,
        )
        async for chunk in stream:
            if chunk.choices and chunk.choices[0].delta.content:
                data = json.dumps({"content": chunk.choices[0].delta.content})
                yield f"data: {data}\n\n"
        yield "data: [DONE]\n\n"

    return StreamingResponse(generate(), media_type="text/event-stream")


# ===== File Attachment Upload =====

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 200_000

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".log", ".csv", ".tsv", ".json", ".xml",
    ".yaml", ".yml", ".toml", ".ini", ".html", ".htm", ".js", ".ts",
    ".py", ".go", ".rs", ".java", ".c", ".cpp", ".h", ".sh", ".sql",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}


def _ext(name: str) -> str:
    if "." not in name:
        return ""
    return "." + name.rsplit(".", 1)[-1].lower()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[PDF extraction requires pypdf. Run python -m pip install -e .]"
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n\n".join(pages).strip()
    except Exception as exc:
        return f"[Failed to parse PDF: {exc}]"


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return "[DOCX extraction requires python-docx. Run python -m pip install -e .]"
    try:
        doc = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()
    except Exception as exc:
        return f"[Failed to parse DOCX: {exc}]"


def _extract_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[XLSX extraction requires openpyxl. Run python -m pip install -e .]"
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: List[str] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows.append(f"## Sheet: {sheet_name}")
            for row in worksheet.iter_rows(max_rows=200, values_only=True):
                rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
        return "\n".join(rows)
    except Exception as exc:
        return f"[Failed to parse XLSX: {exc}]"


def _extract_zip(data: bytes) -> str:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        return f"[Invalid ZIP: {exc}]"

    parts: List[str] = [f"## ZIP contents ({len(archive.namelist())} files)"]
    for name in archive.namelist()[:50]:
        if name.endswith("/"):
            parts.append(f"- {name} (directory)")
            continue
        info = archive.getinfo(name)
        if info.file_size > 1024 * 1024:
            parts.append(f"- {name} ({info.file_size} bytes - skipped)")
            continue
        try:
            inner = archive.read(name)
        except Exception as exc:
            parts.append(f"- {name} (read error: {exc})")
            continue

        ext = _ext(name)
        parts.append(f"\n### {name}")
        if ext in TEXT_EXTS:
            parts.append(inner.decode("utf-8", errors="replace"))
        elif ext == ".pdf":
            parts.append(_extract_pdf(inner))
        elif ext == ".docx":
            parts.append(_extract_docx(inner))
        else:
            parts.append(f"[unsupported file type {ext}, {len(inner)} bytes]")
    return "\n".join(parts)


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    name = file.filename or "uploaded"
    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File too large (max {MAX_UPLOAD_BYTES // (1024 * 1024)} MB)",
        )

    ext = _ext(name)
    if ext in IMAGE_EXTS:
        mime_map = {
            ".png": "png",
            ".jpg": "jpeg",
            ".jpeg": "jpeg",
            ".webp": "webp",
            ".gif": "gif",
            ".bmp": "bmp",
        }
        mime = mime_map.get(ext, "png")
        encoded = base64.b64encode(data).decode("ascii")
        return {
            "kind": "image",
            "filename": name,
            "size": len(data),
            "content": f"data:image/{mime};base64,{encoded}",
        }

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext in (".xlsx", ".xlsm"):
        text = _extract_xlsx(data)
    elif ext == ".zip":
        text = _extract_zip(data)
    elif ext in TEXT_EXTS or not ext:
        text = data.decode("utf-8", errors="replace")
    else:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(status_code=415, detail=f"Unsupported file type: {ext or 'unknown'}")

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + f"\n\n[... truncated, original was {len(text)} chars]"

    return {
        "kind": "text",
        "filename": name,
        "size": len(data),
        "content": text,
    }


app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/")
async def root():
    return FileResponse("static/index.html")
